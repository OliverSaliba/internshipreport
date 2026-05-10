"""
Fill the CSC 480 Word template with ACHI internship content.
Reads: user's template path (default: Downloads).
Writes: ACHI_CSC480_Internship_Report.docx in repo root.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph
from docx.oxml import OxmlElement


def insert_paragraph_after(paragraph: Paragraph, text: str, style=None) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._element.addnext(new_p)
    new_para = Paragraph(new_p, paragraph._parent)
    if text:
        new_para.add_run(text)
    if style is not None:
        new_para.style = style
    return new_para


def set_paragraph_text(paragraph: Paragraph, text: str) -> None:
    paragraph.text = text


def main() -> None:
    repo_root = Path(__file__).resolve().parents[1]
    template_path = Path(r"c:\Users\owner\Downloads\Internship Report Template.docx")
    out_path = repo_root / "ACHI_CSC480_Internship_Report.docx"

    doc = Document(str(template_path))

    par1 = doc.styles["+Par1"]

    # Title page (fix mojibake from template extraction if present)
    set_paragraph_text(doc.paragraphs[0], "Notre Dame University Louaize")
    set_paragraph_text(doc.paragraphs[4], "CSC 480 Internship")
    set_paragraph_text(
        doc.paragraphs[6],
        "Student's Name: [type your full name here]",
    )
    set_paragraph_text(
        doc.paragraphs[7],
        "ID: [type your student ID here]",
    )

    # Preface: remove rubric lines and instructor-only reminders
    set_paragraph_text(doc.paragraphs[12], "")
    set_paragraph_text(
        doc.paragraphs[13],
        "This report describes my internship work on the public website for ACHI Scaffolding, "
        "which promotes engineered scaffolding and temporary works for construction and industrial projects. "
        "Most of my time went into the React application that powers https://achiscaffolding.com/, "
        "including marketing pages, multilingual content, forms, and technical SEO.",
    )
    set_paragraph_text(
        doc.paragraphs[14],
        "The site is a single page application built with Create React App. "
        "I contributed to page layouts, routing and language handling, project and services content, "
        "Italian market pages when that part of the product was enabled, careful attention to accessibility, and the automation that builds "
        "and deploys the static production bundle.",
    )
    set_paragraph_text(doc.paragraphs[16], "")
    set_paragraph_text(
        doc.paragraphs[17],
        "I would like to thank my internship supervisor and the ACHI team for guiding the priorities on the site. "
        "I also thank Notre Dame University Louaize and the Computer Science Department for supporting this internship.",
    )
    for idx in range(19, 24):
        set_paragraph_text(doc.paragraphs[idx], "")

    # Lists
    set_paragraph_text(
        doc.paragraphs[58],
        "Table 2.1: Software categories used during the internship\t2",
    )
    set_paragraph_text(
        doc.paragraphs[64],
        "Figure 3.1: Route structure for the marketing website\t3",
    )

    # Chapter titles (template used a leading colon only)
    set_paragraph_text(doc.paragraphs[67], "Chapter 1: Introduction")
    set_paragraph_text(
        doc.paragraphs[69],
        "ACHI Scaffolding is a construction access company that presents its services online through a marketing "
        "website with contact paths for prospective clients. "
        "The codebase I worked in focuses on the customer facing experience: company story, products, finished "
        "projects, sectors, services, gallery, blog style articles, careers information, and privacy policy. "
        "The interface supports several languages through translation files and URL aware routing.",
    )

    set_paragraph_text(
        doc.paragraphs[71],
        "My objectives were to strengthen the live website in ways that match real business needs. "
        "That meant reliable navigation, clear presentation of projects and services, working contact flows, "
        "and improvements that help search engines understand the pages.",
    )
    set_paragraph_text(
        doc.paragraphs[72],
        "Practical goals included keeping Italian routes and copy coherent when the Italy feature set is enabled, "
        "keeping structured data and meta information aligned with each route, and supporting deployment through "
        "the repository build pipeline.",
    )

    set_paragraph_text(
        doc.paragraphs[74],
        "I expected to work inside a production style front end codebase rather than a classroom demo. "
        "I wanted practice with React Router, translations, form validation, performance conscious media usage, "
        "and build automation on GitHub.",
    )

    set_paragraph_text(
        doc.paragraphs[76],
        "This report has four chapters. Chapter 2 summarizes the software environment. "
        "Chapter 3 explains the work I performed on the ACHI site. Chapter 4 reflects on what I learned.",
    )

    set_paragraph_text(doc.paragraphs[81], "Chapter 2: Technology Background")

    set_paragraph_text(
        doc.paragraphs[83],
        "Development was done on Microsoft Windows using Node.js and npm for installing packages and running scripts. "
        "The front end project lives under a Frontend folder and uses Create React App with react scripts for bundling.",
    )
    set_paragraph_text(
        doc.paragraphs[84],
        "The application relies on React 18 and React Router for client side navigation. "
        "Translations use i18next and react i18next with JSON language files. "
        "Styling combines Tailwind utility classes with Material UI components where helpful. "
        "Forms use Formik with Yup schemas. Motion effects use Framer Motion with attention to reduced motion "
        "preferences on several screens. "
        "SEO related work uses react helmet async, JSON LD components, URL normalization helpers, "
        "and a Node script that generates sitemap output before production builds. "
        "A Three.js based viewer supports optional three dimensional product exploration for selected catalog items. "
        "Static hosting deployment is automated with GitHub Actions that install dependencies, run the production build, "
        "and publish the build folder to GitHub Pages.",
    )

    # Table 2.1
    tbl = doc.tables[0]
    tbl.cell(0, 0).text = "Category"
    tbl.cell(0, 1).text = "What I used on this project"
    tbl.cell(1, 0).text = "Operating system"
    tbl.cell(1, 1).text = "Windows for daily development"
    tbl.cell(2, 0).text = "Core web stack"
    tbl.cell(2, 1).text = "JavaScript with React 18, React Router, Create React App"
    tbl.cell(3, 0).text = "Libraries and tooling"
    tbl.cell(3, 1).text = (
        "i18next, Formik and Yup, Tailwind CSS, Material UI, Framer Motion, react helmet async, Three.js, "
        "sitemap and asset helper scripts, GitHub Actions"
    )

    set_paragraph_text(doc.paragraphs[85], "Table 2.1: Software categories used during the internship")

    set_paragraph_text(
        doc.paragraphs[87],
        "No special laboratory hardware was required. "
        "A standard workstation with a modern browser was enough to develop and manually verify pages.",
    )
    set_paragraph_text(doc.paragraphs[88], "")

    set_paragraph_text(doc.paragraphs[90], "Chapter 3: Description of the Work")

    set_paragraph_text(
        doc.paragraphs[92],
        "My contributions centered on the React routes and components that visitors actually see. "
        "The route table connects paths such as home, about, products (including a dedicated three dimensional "
        "product view), projects and individual project details, sectors, services and single service pages, "
        "gallery, blog entries, careers pages, and privacy policy. "
        "Language routing wraps those routes so URLs and selected language stay consistent. "
        "When the Italy feature flag is enabled, Italian experiences include dedicated routes such as "
        "settori applicazioni plus Italian copy stored in translation data.",
    )
    set_paragraph_text(
        doc.paragraphs[93],
        "The routing flow can be described in simple terms. The browser requests a path, React Router selects "
        "a page component, and that page pulls translated strings, images from the public assets, and SEO metadata.",
    )
    set_paragraph_text(doc.paragraphs[94], "Figure 3.1: Route structure for the marketing website")
    set_paragraph_text(doc.paragraphs[95], "")
    set_paragraph_text(
        doc.paragraphs[96],
        "The following subsections follow the course outline for describing internship tasks.",
    )

    problems_para_a = doc.paragraphs[107]
    problems_para_b = doc.paragraphs[108]
    impact_para = doc.paragraphs[110]
    spare_heading_para = doc.paragraphs[112]
    conclusion_heading_para = doc.paragraphs[113]
    conclusion_para_a = doc.paragraphs[115]
    conclusion_para_b = doc.paragraphs[116]
    conclusion_para_c = doc.paragraphs[117]
    self_eval_para = doc.paragraphs[119]
    tail_para_a = doc.paragraphs[120]
    tail_para_b = doc.paragraphs[121]

    subsection_bodies = [
        (
            "Requirements Analysis",
            "Requirements came from business priorities on the live site. "
            "Examples included presenting projects with correct imagery per project id, keeping Italian pages "
            "complete where required, ensuring contact submissions validate before send, and keeping SEO "
            "elements truthful and aligned with visible content. "
            "Because text lives in translation JSON, changes had to preserve keys and structure so pages "
            "would not break at runtime.",
        ),
        (
            "Software Design",
            "The project favors reusable UI pieces such as header, footer, cards, and SEO helpers, "
            "while each major URL maps to a page level component. "
            "Feature flags isolate market specific routes. "
            "Structured business information for search engines is centralized so it stays consistent "
            "across the shell of the application.",
        ),
        (
            "User Interface Design",
            "Pages combine Tailwind layout utilities with occasional CSS modules for distinctive sections. "
            "Carousels and motion are used for marketing emphasis but I learned to test scrolling and focus behavior "
            "so the site stayed usable. "
            "Imagery is prominent on projects and gallery screens, so lazy loading and sensible formats matter for speed.",
        ),
        (
            "Coding",
            "Coding work meant editing JSX components, adjusting translation files, wiring navigation helpers, "
            "and fixing bugs where media paths or language aware links did not match the intended project record. "
            "I also touched SEO utilities such as route level metadata patterns and guards related to secure content.",
        ),
        (
            "Testing",
            "The repository includes Create React App tests with Testing Library, but most verification during "
            "internship tasks was manual browsing across routes, languages, and form edge cases. "
            "Build verification included confirming generated sitemap output and that the hosting workflow "
            "packaged the expected static files.",
        ),
        (
            "Documentation",
            "Several markdown notes in the repository explain SEO decisions and hosting assumptions. "
            "README files still carry some Create React App boilerplate that should stay conflict free for teammates.",
        ),
        (
            "Preparing Training Manuals",
            "I did not prepare a formal training manual. "
            "Knowledge transfer would rely on the README, environment variables such as the contact script URL, "
            "and short explanations inside pull requests.",
        ),
        (
            "Deployment (Implementation, Training, etc.)",
            "Local development uses npm start. "
            "Production builds run sitemap generation first, then react scripts build, followed by a small script "
            "that copies a custom not found page into the build output. "
            "GitHub Actions installs Frontend dependencies with npm ci, runs the build with CI friendly settings, "
            "and publishes the build directory as a Pages artifact.",
        ),
        (
            "Customer Support",
            "The website supports ordinary commercial contact through visible phone and email details plus forms. "
            "Italian pages include consent language where relevant. "
            "My work indirectly supports customer support by keeping those entry points accurate and trustworthy.",
        ),
    ]

    for title, body in reversed(subsection_bodies):
        for p in doc.paragraphs:
            if p.text.strip() == title:
                insert_paragraph_after(p, body, par1)
                break

    set_paragraph_text(
        problems_para_a,
        "One recurring challenge was separating routing bugs from content bugs. "
        "A wrong image or paragraph often traced back to translation data or a media map rather than the router itself. "
        "Another challenge was keeping SEO changes disciplined so titles and canonical patterns stayed coherent "
        "when URLs moved slightly.",
    )
    set_paragraph_text(
        problems_para_b,
        "From an ethics standpoint, a public company site must avoid overstated claims and must handle personal "
        "data from forms responsibly. "
        "Mixed content and insecure scripts can undermine trust, so defensive initialization around secure loading "
        "is part of professional responsibility, not only a technical detail.",
    )

    set_paragraph_text(
        impact_para,
        "Clear online presentation supports ACHI in Lebanon and in broader markets where prospects compare vendors "
        "before they call. "
        "Better structured data and cleaner routing help people discover services through search. "
        "Reliable forms shorten the distance between interest and a real conversation.",
    )

    set_paragraph_text(spare_heading_para, "")
    set_paragraph_text(conclusion_heading_para, "Chapter 4: Conclusion")

    set_paragraph_text(
        conclusion_para_a,
        "I gained concrete experience shipping UI changes in a React codebase that already serves real traffic. "
        "I practiced tracing issues from what the user sees down to data files and configuration.",
    )
    set_paragraph_text(
        conclusion_para_b,
        "I became more comfortable with multilingual routing, pragmatic SEO, and the rhythm of build and deploy steps "
        "that surround application code.",
    )
    set_paragraph_text(
        conclusion_para_c,
        "The internship reminded me that finishing a small visible detail on a marketing site can matter to "
        "a business conversation even when the internal diff looks minor.",
    )

    set_paragraph_text(
        self_eval_para,
        "What went well for me was persistence when debugging mapping mistakes between projects and assets, "
        "and checking pages in more than one language before declaring the task done. "
        "What I would improve is adding more repeatable automated checks for critical routes so regressions "
        "are caught earlier. "
        "I also want to keep refining how I document assumptions for teammates who maintain hosting secrets "
        "and third party scripts.",
    )
    set_paragraph_text(tail_para_a, "")
    set_paragraph_text(tail_para_b, "")

    doc.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
