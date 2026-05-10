"""
Generate CSC 480 internship report content from ACHI- repo facts.
Output: Internship_Report_ACHI_Filled.docx (repo root)
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.shared import Pt


def set_body_style(p, size=12):
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    p.paragraph_format.line_spacing = 1.5
    for run in p.runs:
        run.font.size = Pt(size)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        run.font.size = Pt(14 if level == 1 else 13)
    return p


def add_para(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(12)
    set_body_style(p)
    return p


def main():
    root = Path(__file__).resolve().parents[1]
    out = root / "Internship_Report_ACHI_Filled.docx"

    doc = Document()

    # Title block
    t = doc.add_paragraph()
    t.alignment = 1  # center
    for line, bold, sz in [
        ("Notre Dame University – Louaize", True, 14),
        ("Faculty of Natural and Applied Sciences", False, 12),
        ("Computer Science Department", False, 12),
        ("CSC 480 – Internship", True, 12),
        ("Final Report", True, 14),
        ("", False, 12),
        ("Student’s Name: [YOUR FULL NAME]", False, 12),
        ("ID: [YOUR STUDENT ID]", False, 12),
        ("Fall 2025", False, 12),
    ]:
        if not line:
            doc.add_paragraph()
            continue
        p = doc.add_paragraph()
        p.alignment = 1
        r = p.add_run(line)
        r.bold = bold
        r.font.size = Pt(sz)

    doc.add_page_break()

    add_heading(doc, "Preface and Acknowledgment", 1)
    add_para(
        doc,
        "This internship report documents work performed on the public-facing web application for ACHI Scaffolding "
        "(https://achiscaffolding.com/), a company that provides engineered scaffolding, temporary works, and related "
        "services. The internship focused on front-end development of a multilingual marketing and lead-generation "
        "website implemented as a single-page application (SPA) using React.",
    )
    add_para(
        doc,
        "The main deliverables included implementing and refining user-facing pages (home, projects, services, "
        "contact, blog, gallery, product catalog with 3D preview), improving Italian-market content and routing, "
        "SEO and structured data (JSON-LD, sitemap generation), accessibility-oriented UI patterns, and integration "
        "of contact flows (Formik/Yup validation, external script submission).",
    )
    add_para(
        doc,
        "I acknowledge the supervision and feedback from my internship supervisor and the ACHI team. "
        "Finally, I thank Notre Dame University – Louaize and the Computer Science Department for the opportunity "
        "to apply coursework in a professional software project.",
    )

    doc.add_page_break()
    add_heading(doc, "Table of Contents", 1)
    toc = """Preface and Acknowledgment
Chapter 1: Introduction
  1.1 Company Background
  1.2 Internship Objectives
  1.3 My Expectations
  1.4 Structure of the Report
Chapter 2: Technology Background
  2.1 Software
  2.2 Hardware
Chapter 3: Description of the Work
  3.1 Work Done as an Intern
  3.2 Problems Encountered
  3.3 Impact of Your Work
Chapter 4: Conclusion
  4.1 Experience Gained During the Internship
  4.2 Self-Evaluation"""
    for line in toc.split("\n"):
        add_para(doc, line)

    doc.add_page_break()
    add_heading(doc, "Chapter 1: Introduction", 1)

    add_heading(doc, "1.1 Company Background", 2)
    add_para(
        doc,
        "ACHI Scaffolding (branded online as ACHI) operates in the construction access sector: scaffolding systems, "
        "façade access, temporary structures, and related services. The production website targets Lebanon and "
        "broader regions, with additional Italian-language experiences and Italy-specific routes (for example "
        "Italian-only pages such as “Settori di applicazione” and Italian project case studies) controlled by "
        "feature flags and language-aware routing in the codebase.",
    )
    add_para(
        doc,
        "The internship repository centers on the front-end application under Frontend/, built with Create React App "
        "(react-scripts 5). The application homepage is configured as https://achiscaffolding.com/ in package.json.",
    )

    add_heading(doc, "1.2 Internship Objectives", 2)
    add_para(
        doc,
        "Primary objectives aligned with the project were: (1) implement and maintain React components and pages "
        "with consistent UX; (2) support internationalization (i18next) for multiple languages; (3) implement forms "
        "with validation and safe submission patterns; (4) improve SEO artifacts (meta tags, canonical URLs, "
        "sitemap generation scripts); (5) collaborate on content-heavy sections (projects, services, AEO/FAQ blocks) "
        "without breaking routing or translation keys.",
    )

    add_heading(doc, "1.3 My Expectations", 2)
    add_para(
        doc,
        "I expected to gain experience with a real deployed React codebase, Git-based collaboration, and the full "
        "cycle from UI implementation to production build. I also expected exposure to performance and SEO concerns "
        "typical of marketing sites (image formats, lazy loading, structured data).",
    )

    add_heading(doc, "1.4 Structure of the Report", 2)
    add_para(
        doc,
        "Chapter 2 summarizes the software stack and tooling inferred from package.json and project structure. "
        "Chapter 3 describes concrete modules and features implemented or modified in the repository. "
        "Chapter 4 reflects on outcomes and self-evaluation.",
    )

    doc.add_page_break()
    add_heading(doc, "Chapter 2: Technology Background", 1)

    add_heading(doc, "2.1 Software", 2)
    add_para(
        doc,
        "Operating system: Windows 10/11 (development environment on the intern workstation). "
        "Runtime: Node.js with npm for dependency management.",
    )
    add_para(
        doc,
        "Core framework and libraries (from Frontend/package.json): React 18.2, react-dom, react-router-dom 6.17 "
        "for client-side routing, react-i18next and i18next for translations, framer-motion for UI motion, "
        "Formik and Yup for forms, @mui/material and @emotion for Material UI components, Tailwind CSS 3.2 as a "
        "utility-first styling layer, react-slick and slick-carousel for carousels, three for optional 3D product "
        "visualization (Product3DView), react-helmet-async for document head management, and testing utilities "
        "(@testing-library/react). Build tooling is provided by Create React App (react-scripts 5.0.1).",
    )
    add_para(
        doc,
        "Additional project scripts include generate:sitemap (node scripts/generateSitemap.js) invoked before build, "
        "postbuild copy of a 404 page (scripts/copy-404.js), and optional WebP generation (scripts/generateWebp.js). "
        "Deployment-related tooling includes gh-pages as a devDependency for static hosting workflows.",
    )
    add_para(
        doc,
        "Table 2.1 – Primary software stack (summary)",
        bold=True,
    )
    add_para(
        doc,
        "Category | Technology\n"
        "Language | JavaScript (React)\n"
        "UI | React components, Tailwind CSS, MUI\n"
        "Routing | React Router v6\n"
        "i18n | i18next / react-i18next\n"
        "Forms | Formik + Yup\n"
        "Motion | Framer Motion\n"
        "SEO | Helmet, JSON-LD components, sitemap script\n"
        "Build | Create React App (webpack/babel via react-scripts)",
    )

    add_heading(doc, "2.2 Hardware", 2)
    add_para(
        doc,
        "No specialized industrial hardware was required for this internship. Development used standard office "
        "hardware: a PC with sufficient RAM for Node/Webpack builds, a display for UI testing, and network access "
        "for dependency installation and deployment verification.",
    )

    doc.add_page_break()
    add_heading(doc, "Chapter 3: Description of the Work", 1)

    add_heading(doc, "3.1 Work Done as an Intern", 2)
    add_para(
        doc,
        "The repository implements a marketing and lead-generation website with many routed pages. Key areas "
        "observed and worked on in the codebase include:",
    )
    add_para(
        doc,
        "Routing and layout: Frontend/src/routes/AppRoutes.js defines the main route tree (Home, About, Products, "
        "Product3DView, Projects, Sectors, SettoriApplicazioniIt when Italy is enabled, ProjectDetails, Gallery, "
        "Blog, BlogItem, Services, SingleService, Careers, PrivacyPolicy, PageNotFound). Language handling is "
        "integrated via LangRouter and URL language segments.",
    )
    add_para(
        doc,
        "Projects experience: Projects.js lists featured projects for Italian vs default locales; ProjectDetails.js "
        "loads per-project media (hero and gallery images) and binds long-form Italian content from translation JSON "
        "(it.json under projectDetails.items.*). ProjectsOverview.js renders the “projects” strip on the home page "
        "with SmartLink navigation to /project/:id.",
    )
    add_para(
        doc,
        "Italian market pages: SettoriApplicazioniIt.jsx implements the Italian-only “settori” page with hero video, "
        "glass-style hero panel (CSS module), sector cards with Font Awesome icons aligned to shared sector data "
        "(settoriApplicazioniIt.js), and a Formik/Yup contact form posting to a configurable script URL "
        "(REACT_APP_CONTACT_SCRIPT_URL).",
    )
    add_para(
        doc,
        "SEO and quality: components under seo/ (e.g., RouteSeo, sitemap routes), JSON-LD in pages like Home.js and "
        "About.js, and defensive helpers such as mixed content guard initialization in AppRoutes.js.",
    )

    add_heading(doc, "3.1.1 Requirements Analysis", 3)
    add_para(
        doc,
        "Requirements were inferred from existing routes, translation keys, and stakeholder requests (e.g., correct "
        "image mapping per project slug, Italian-only flows, hero/gallery consistency). The codebase separates "
        "content (JSON translations) from presentation (React components), which constrained changes to avoid "
        "breaking i18n keys.",
    )

    add_heading(doc, "3.1.2 Software Design", 3)
    add_para(
        doc,
        "The design follows a component-based architecture: reusable UI (Header, Footer, ContactForm), page-level "
        "containers, and SEO components. Feature flags (e.g., ITALY_ENABLED) gate Italian-only routes and UI.",
    )

    add_heading(doc, "3.1.3 User Interface Design", 3)
    add_para(
        doc,
        "UI work combined Tailwind utility classes with occasional CSS modules for complex sections (e.g., Italian "
        "hero glass panels, sectors bar styling). Motion was applied via framer-motion for scroll reveals and hero "
        "animations, with reduced-motion awareness in several pages.",
    )

    add_heading(doc, "3.1.4 Coding", 3)
    add_para(
        doc,
        "Coding tasks included JSX/React updates, translation JSON edits for Italian copy, image path corrections in "
        "ProjectsOverview.js, Projects.js, and ProjectDetails.js (PROJECT_MEDIA map), and styling adjustments for "
        "readability and brand colors (#28509E accent appears across UI).",
    )

    add_heading(doc, "3.1.5 Testing", 3)
    add_para(
        doc,
        "The project includes CRA test scaffolding (@testing-library). Practical testing during internship work "
        "centered on manual route verification, visual regression checks for hero/gallery swaps, and form "
        "validation behavior in browsers.",
    )

    add_heading(doc, "3.1.6 Documentation", 3)
    add_para(
        doc,
        "README.md in Frontend/ retains Create React App boilerplate (note: resolve any merge markers before "
        "submission). SEO documentation is embedded in code via seoConfig and sitemap route lists.",
    )

    add_heading(doc, "3.1.7 Preparing Training Manuals", 3)
    add_para(
        doc,
        "Not a primary deliverable in the repository; any internal handover would be via README/runbook and "
        "environment variables documented in code (e.g., REACT_APP_CONTACT_SCRIPT_URL).",
    )

    add_heading(doc, "3.1.8 Deployment", 3)
    add_para(
        doc,
        "Production build pipeline: npm run build runs sitemap generation then react-scripts build; postbuild copies "
        "404 handling. The configured homepage is https://achiscaffolding.com/.",
    )

    add_heading(doc, "3.1.9 Customer Support", 3)
    add_para(
        doc,
        "The site exposes contact channels (phone, WhatsApp, contact forms) consistent with business schema in "
        "AppRoutes.js (+96103322811, achi.gr@hotmail.com) and Italian-specific CTAs where translations provide them.",
    )

    add_para(
        doc,
        "Figure 3.1 (placeholder): High-level route diagram — Browser → React Router → Page components → i18n JSON / "
        "public assets → SEO head + JSON-LD.",
    )

    add_heading(doc, "3.2 Problems Encountered", 2)
    add_para(
        doc,
        "Challenges included: (1) separating routing correctness from content and media mapping errors when multiple "
        "locales share similar project names; (2) maintaining SEO and translation key integrity while swapping images; "
        "(3) balancing decorative motion with accessibility (prefers-reduced-motion patterns in parts of the app); "
        "(4) merge-conflict residue in README that should be cleaned for professional repository hygiene.",
    )
    add_para(
        doc,
        "Ethical/legal/social considerations: public websites must present accurate project imagery and truthful "
        "service descriptions; forms must respect privacy expectations (GDPR copy on Italian forms) and avoid broken "
        "analytics or mixed-content issues (mixedContentGuard).",
    )

    add_heading(doc, "3.3 Impact of Your Work", 2)
    add_para(
        doc,
        "Improved clarity of Italian project presentation supports trust for EU-based prospects. SEO and structured "
        "data improve discoverability. Stable forms and CTAs reduce friction for commercial inquiries, with global "
        "reach via multilingual content.",
    )

    doc.add_page_break()
    add_heading(doc, "Chapter 4: Conclusion", 1)

    add_heading(doc, "4.1 Experience Gained During the Internship", 2)
    add_para(
        doc,
        "I strengthened skills in React, component composition, router-based localization, and pragmatic SEO for "
        "marketing sites. I learned to navigate a mid-size front-end codebase, trace data from translation files to "
        "UI, and validate changes across multiple pages that reuse similar patterns.",
    )

    add_heading(doc, "4.2 Self-Evaluation", 2)
    add_para(
        doc,
        "Strengths: attention to cross-page consistency (home cards vs detail media), willingness to iterate on UI "
        "feedback, and systematic tracing of bugs to data sources rather than only components. "
        "Areas to improve: earlier verification of asset-to-project mapping with stakeholders, automated visual or "
        "route tests for critical marketing pages, and keeping repository documentation conflict-free.",
    )

    add_para(
        doc,
        "Replace bracketed placeholders ([YOUR FULL NAME], [YOUR STUDENT ID]) and adjust dates/duration with your "
        "actual internship contract before submission.",
    )

    doc.save(out)
    print(f"Wrote: {out}")


if __name__ == "__main__":
    main()
